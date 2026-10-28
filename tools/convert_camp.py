"""Extractor for Camp Structures & Props from '영덱 공략 종합_전체 야영지.xlsx'
and Growth assets.

Generates:
  - data/camp.json & assets/camp/*.webp
  - assets/growth/underdark_growth.png
  - updates data/underdark_all.js
"""
from __future__ import annotations

import io
import json
import re
import shutil
from pathlib import Path
from PIL import Image
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = ROOT / '전체 야영지.xlsx'
if not XLSX_PATH.exists():
    XLSX_PATH = ROOT / '영덱 공략 종합_전체 야영지.xlsx'
GROWTH_SRC = ROOT / '언더다크 성장.png'
DATA_DIR = ROOT / 'data'
IMG_CAMP = ROOT / 'assets' / 'camp'
IMG_GROWTH = ROOT / 'assets' / 'growth'

CHO = 'ㄱㄲㄴㄷㄸㄹㅁㅂㅃㅅㅆㅇㅈㅉㅊㅋㅌㅍㅎ'


def chosung(s: str) -> str:
    out = []
    for ch in s:
        code = ord(ch) - 0xAC00
        if 0 <= code < 11172:
            out.append(CHO[code // 588])
        elif ch.strip():
            out.append(ch.lower())
    return ''.join(out)


def clean(s: str) -> str:
    return re.sub(r'\s+', ' ', (s or '').replace('\n', ' ')).strip()


def norm_key(s: str) -> str:
    return re.sub(r'[\s/·,.\-()★☆\[\]]+', '', s).lower()


def main():
    IMG_CAMP.mkdir(parents=True, exist_ok=True)
    IMG_GROWTH.mkdir(parents=True, exist_ok=True)

    # 1. Copy Growth image
    if GROWTH_SRC.exists():
        growth_dst = IMG_GROWTH / 'underdark_growth.png'
        shutil.copy2(GROWTH_SRC, growth_dst)
        print(f"Copied {GROWTH_SRC.name} -> {growth_dst.relative_to(ROOT)}")

    # 2. Open workbook
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb['전체 야영지']

    # Clean old camp images so renumbered items don't leave stale files
    for old_f in IMG_CAMP.glob('*.webp'):
        old_f.unlink()

    # Map images by 0-based anchor row
    # Image in col 3 (D열) in updated sheet, or col 2 (C열) in legacy sheet
    images_by_row = {}
    for img in ws._images:
        r = img.anchor._from.row
        c = img.anchor._from.col
        if c in (2, 3):
            images_by_row[r] = img

    def save_camp_image(img, out_name: str) -> str:
        out_path = IMG_CAMP / out_name
        im_data = img._data()
        im = Image.open(io.BytesIO(im_data)).convert('RGBA')
        im.thumbnail((180, 180), Image.LANCZOS)
        im.save(out_path, 'WEBP', quality=90, method=6)
        return out_path.relative_to(ROOT).as_posix()

    def parse_level(val) -> int:
        if val is None:
            return 0
        try:
            return int(val)
        except (ValueError, TypeError):
            m = re.search(r'\d+', str(val))
            return int(m.group(0)) if m else 0

    # 3. Parse Structures (Rows 7 ~ 16, 0-based 6 ~ 15)
    structures = []
    for r in range(7, 17):
        name = clean(ws.cell(r, 2).value)
        level = parse_level(ws.cell(r, 3).value)
        abil = clean(ws.cell(r, 5).value)
        if not name:
            continue

        idx = len(structures) + 1  # 1 to 10
        img_rel = None
        img_obj = images_by_row.get(r - 1)
        if img_obj:
            img_rel = save_camp_image(img_obj, f"structure_{idx:02d}_{norm_key(name)}.webp")

        structures.append({
            'id': f"structure-{idx}",
            'no': idx,
            'type': 'structure',
            'typeName': '구조물',
            'name': name,
            'level': level,
            'ability': abil,
            'img': img_rel,
            'search': {
                'name': norm_key(name),
                'cho': chosung(name),
                'text': norm_key(f"{name} {level}레벨 {level}렙 {abil}"),
            }
        })

    # 4. Parse Props (Rows 20 ~ 53, 0-based 19 ~ 52)
    props = []
    for r in range(20, 54):
        name = clean(ws.cell(r, 2).value)
        level = parse_level(ws.cell(r, 3).value)
        abil = clean(ws.cell(r, 5).value)
        if not name:
            continue

        idx = len(props) + 1  # 1 to 34
        img_rel = None
        img_obj = images_by_row.get(r - 1)
        if img_obj:
            img_rel = save_camp_image(img_obj, f"prop_{idx:02d}_{norm_key(name)}.webp")

        props.append({
            'id': f"prop-{idx}",
            'no': idx,
            'type': 'prop',
            'typeName': '소품',
            'name': name,
            'level': level,
            'ability': abil,
            'img': img_rel,
            'search': {
                'name': norm_key(name),
                'cho': chosung(name),
                'text': norm_key(f"{name} {level}레벨 {level}렙 {abil}"),
            }
        })

    # 5. Common info: Grades, Sub-options, Random options (Col 7 / G열)
    grades = []
    for r in range(7, 15):
        g = clean(ws.cell(r, 7).value)
        if g:
            grades.append(g)

    sub_options = []
    for r in range(20, 28):
        opt = clean(ws.cell(r, 7).value)
        if opt:
            sub_options.append(opt)

    random_options = []
    for r in range(30, 42):
        opt = clean(ws.cell(r, 7).value)
        if opt:
            random_options.append(opt)

    level_note = clean(ws.cell(4, 2).value) or '레벨 - 계정 레벨에 따라 구조물과 소품의 획득이 가능해집니다.'

    camp_data = {
        'title': '언더다크 : 야영지 구조물 / 소품',
        'updatedAt': '2026-10-28',
        'author': 'Ara-달품',
        'common': {
            'levelNote': level_note,
            'grades': grades,
            'subOptions': sub_options,
            'randomOptions': random_options,
        },
        'structures': structures,
        'props': props,
    }

    # Save data/camp.json
    out_json = DATA_DIR / 'camp.json'
    with open(out_json, 'w', encoding='utf-8') as f:
        json.dump(camp_data, f, ensure_ascii=False, indent=2)
    print(f"Saved {out_json.relative_to(ROOT)} ({len(structures)} structures, {len(props)} props)")

    # Update underdark_all.js
    all_js_path = DATA_DIR / 'underdark_all.js'
    if all_js_path.exists():
        with open(all_js_path, 'r', encoding='utf-8') as f:
            content = f.read()

        prefix = "window.UNDERDARK_EXPANSION="
        if content.startswith(prefix):
            raw = content[len(prefix):].rstrip(';')
            expansion = json.loads(raw)
            expansion['camp'] = camp_data
            with open(all_js_path, 'w', encoding='utf-8') as f:
                f.write(prefix + json.dumps(expansion, ensure_ascii=False, separators=(',', ':')) + ';')
            print("Updated data/underdark_all.js with camp data.")

if __name__ == '__main__':
    main()
