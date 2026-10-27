import json
import re
from pathlib import Path
import openpyxl
import sys

sys.path.insert(0, r"c:\dev\workplace\underdark_web\tools")
from stats import extract_stats, STAT_DEFS

ROOT = Path(r"c:\dev\workplace\underdark_web")

print("==================================================")
print("     UNDERDARK WEB COMPREHENSIVE DATA AUDIT       ")
print("==================================================")

# 1. Audit Items Master Sheet vs items.json
wb = openpyxl.load_workbook(ROOT / 'Ara_DB_구글시트_마스터.xlsx', data_only=True)
ws_items = wb['아이템']
headers = [cell.value for cell in ws_items[1]]
col_map = {h: i for i, h in enumerate(headers)}

with open(ROOT / 'data' / 'items.json', encoding='utf-8') as f:
    items_json_data = json.load(f)
    items_json = items_json_data['items']

items_map = {it['name']: it for it in items_json}

item_issues = []
checked_items = 0

for row in ws_items.iter_rows(min_row=2, values_only=True):
    if not row or not any(row): continue
    name = row[col_map['아이템명']]
    if not name: continue
    checked_items += 1
    
    it = items_map.get(name)
    if not it:
        item_issues.append(f"[MISSING ITEM] '{name}' in sheet but not in items.json")
        continue
    
    # Check base ability text
    sheet_base = (row[col_map['기본 고유 능력']] or '').strip()
    json_base = '\n'.join(it.get('abilities', []))
    # normalise spaces
    s_norm = re.sub(r'\s+', ' ', sheet_base)
    j_norm = re.sub(r'\s+', ' ', json_base)
    if s_norm != j_norm:
        item_issues.append(f"[BASE ABILITY DIFF] {name}:\n   Sheet: {s_norm}\n   JSON : {j_norm}")
    
    # Check awakening text
    for lv in range(1, 6):
        col = f'{lv}각성 능력'
        sheet_awak = (row[col_map.get(col)] or '').strip() if col in col_map else ''
        json_awak = '\n'.join(it.get('awakening', {}).get(str(lv), []))
        s_norm = re.sub(r'\s+', ' ', sheet_awak)
        j_norm = re.sub(r'\s+', ' ', json_awak)
        if s_norm != j_norm:
            item_issues.append(f"[AWAKENING DIFF] {name} (Lv{lv}):\n   Sheet: {s_norm}\n   JSON : {j_norm}")

    # Check stats extraction
    base_stats = extract_stats(sheet_base)
    for k, v in base_stats.items():
        val = it.get('stats', {}).get(k, [None])[0]
        if val != v:
            item_issues.append(f"[STAT MISMATCH] {name} Lv0 {k}: extracted {v} vs json {val}")
            
    for lv in range(1, 6):
        col = f'{lv}각성 능력'
        if col in col_map and row[col_map[col]]:
            txt = row[col_map[col]]
            extracted = extract_stats(txt)
            for k, v in extracted.items():
                stat_list = it.get('stats', {}).get(k, [])
                val = stat_list[lv] if lv < len(stat_list) else None
                if val != v:
                    item_issues.append(f"[STAT MISMATCH] {name} Lv{lv} {k}: extracted {v} vs json {val}")

print(f"Items checked: {checked_items}")
print(f"Item issues found: {len(item_issues)}")
for iss in item_issues[:20]:
    print("  ", iss)

# 2. Audit Heroes Master Sheet vs heroes.json
ws_heroes = wb['영웅']
h_headers = [cell.value for cell in ws_heroes[1]]
h_cols = {h: i for i, h in enumerate(h_headers)}

with open(ROOT / 'data' / 'heroes.json', encoding='utf-8') as f:
    heroes_json = json.load(f)
heroes_map = {h['name']: h for h in heroes_json}

hero_issues = []
h_count = 0
for row in ws_heroes.iter_rows(min_row=2, values_only=True):
    if not row or not any(row): continue
    name = row[h_cols['영웅명']]
    if not name: continue
    h_count += 1
    h = heroes_map.get(name)
    if not h:
        hero_issues.append(f"[MISSING HERO] '{name}'")
        continue
    # check tier
    sheet_tier = row[h_cols.get('티어')]
    if sheet_tier and h.get('tier') != sheet_tier:
        hero_issues.append(f"[HERO TIER DIFF] {name}: sheet={sheet_tier} vs json={h.get('tier')}")

print(f"\nHeroes checked: {h_count}, Issues: {len(hero_issues)}")

# 3. Audit Mercenaries Master Sheet vs mercenaries.json
ws_mercs = wb['지원용병']
m_headers = [cell.value for cell in ws_mercs[1]]
m_cols = {m: i for i, m in enumerate(m_headers)}

with open(ROOT / 'data' / 'mercenaries.json', encoding='utf-8') as f:
    mercs_json = json.load(f)
mercs_map = {m['name']: m for m in mercs_json}

merc_issues = []
m_count = 0
for row in ws_mercs.iter_rows(min_row=2, values_only=True):
    if not row or not any(row): continue
    name = row[m_cols['용병명']]
    if not name: continue
    m_count += 1
    m = mercs_map.get(name)
    if not m:
        merc_issues.append(f"[MISSING MERC] '{name}'")
        continue
    sheet_rec = row[m_cols.get('추천 등급')]
    if sheet_rec and m.get('recommend') != sheet_rec:
        merc_issues.append(f"[MERC REC DIFF] {name}: sheet={sheet_rec} vs json={m.get('recommend')}")

print(f"Mercenaries checked: {m_count}, Issues: {len(merc_issues)}")

# 4. Audit Mounts Master Sheet vs mounts.json
ws_mounts = wb['탈것']
mt_headers = [cell.value for cell in ws_mounts[1]]
mt_cols = {mt: i for i, mt in enumerate(mt_headers)}

with open(ROOT / 'data' / 'mounts.json', encoding='utf-8') as f:
    mounts_json = json.load(f)['mounts']
mounts_map = {m['name']: m for m in mounts_json}

mount_issues = []
mt_count = 0
for row in ws_mounts.iter_rows(min_row=2, values_only=True):
    if not row or not any(row): continue
    name = row[mt_cols['탈것 이름']]
    if not name: continue
    mt_count += 1
    m = mounts_map.get(name)
    if not m:
        mount_issues.append(f"[MISSING MOUNT] '{name}'")
        continue

print(f"Mounts checked: {mt_count}, Issues: {len(mount_issues)}")

# 5. Audit Pets Master Sheet vs pets.json
ws_pets = wb['펫']
p_headers = [cell.value for cell in ws_pets[1]]
p_cols = {p: i for i, p in enumerate(p_headers)}

with open(ROOT / 'data' / 'pets.json', encoding='utf-8') as f:
    pets_json = json.load(f)['pets']
pets_map = {p['name']: p for p in pets_json}

pet_issues = []
p_count = 0
for row in ws_pets.iter_rows(min_row=2, values_only=True):
    if not row or not any(row): continue
    name = row[p_cols['펫 이름']]
    if not name: continue
    p_count += 1
    p = pets_map.get(name)
    if not p:
        pet_issues.append(f"[MISSING PET] '{name}'")
        continue

print(f"Pets checked: {p_count}, Issues: {len(pet_issues)}")

# 6. Audit Camp Data
with open(ROOT / 'data' / 'camp.json', encoding='utf-8') as f:
    camp_json = json.load(f)
structures = camp_json['structures']
props = camp_json['props']
common = camp_json['common']
print(f"\nCamp Structures: {len(structures)}, Props: {len(props)}")
print(f"Camp Grades: {len(common['grades'])}, Sub Options: {len(common['subOptions'])}, Random Options: {len(common['randomOptions'])}")

# 7. Check All Stat Value Anomalies
print("\n=== Stat Value Sanity Checks across all 208 Items ===")
anomalies = []
for it in items_json:
    for stat, vals in it.get('stats', {}).items():
        for lv, v in enumerate(vals):
            if v is not None:
                if v < 0:
                    anomalies.append(f"Negative stat: {it['name']} Lv{lv} [{stat}] = {v}")
                if v > 1000 and stat not in ('bomb_dmg',):
                    anomalies.append(f"Extremely high stat: {it['name']} Lv{lv} [{stat}] = {v}")

print(f"Stat anomalies found: {len(anomalies)}")
for a in anomalies:
    print("  ", a)

print("\n==================================================")
print("             AUDIT COMPLETE                       ")
print("==================================================")
